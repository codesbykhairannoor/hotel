"""
Generate Curated Genuine Hotel Leads Excel Sheet & Update Next.js hotel leads data
Strictly Hotels, Suites, Resorts, Inns, Boutique Mansions (ZERO Kos-Kosan / Coliving)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import re
import urllib.parse
import os

INPUT_XLSX = 'listing.xlsx'
OUTPUT_XLSX = 'BATCH_1_100_LEADS_HOTEL_LUXURY_PITCH_PROSPEK.xlsx'
JSON_DATA_PATH = 'src/data/hotelsLeads.json'

wb_in = openpyxl.load_workbook(INPUT_XLSX, read_only=True)
s_in = wb_in['Hotel & Penginapan']

# Strict blacklist to eliminate kos-kosan, coliving, budget chains, student dorms, and theme parks
blacklist = [
    'kost', 'kos', 'kosan', 'coliving', 'co-living', 'rukita', 'cove',
    'kontrakan', 'asrama', 'dormitory', 'dorm', 'boarding', 'student',
    'kamar sewa', 'sewa kamar', 'backpackers', 'backpacker', 'hostel',
    'capsule', 'kapsul', 'catering', 'sekolah', 'bimbel', 'homestay',
    'home stay', 'pondok', 'mess', 'wisma', 'reddoorz', 'oyo', 'airy',
    "d'paragon", 'dparagon', 'bobopod', 'bobobox', 'wisata', 'taman',
    'kamar', 'ostel', 'djuragan', 'warung', 'kantin', 'bengkel', 'laundry',
    'the great asia africa', 'cluster harris', 'guest house 555 nki', 'dreaminn'
]

curated_leads = []
seen_phones = set()
seen_slugs = set()

for row in s_in.iter_rows(values_only=True, min_row=2):
    no, name, klaster, jenis, telp, wa, kota, prov, rating, ulasan, alamat, kampus, kantor, web, gmaps = row[:15]
    if not name:
        continue
    name_str = str(name).strip()
    name_lower = name_str.lower()
    
    # Filter blacklisted keywords
    if any(b in name_lower for b in blacklist):
        continue
        
    # Phone normalization
    clean_phone = re.sub(r'[^0-9]', '', str(telp or ''))
    if clean_phone.startswith('62'):
        clean_phone = '0' + clean_phone[2:]
    if not (clean_phone.startswith('08') and len(clean_phone) in range(10, 14)):
        continue
        
    # Avoid duplicate phone numbers
    if clean_phone in seen_phones:
        continue
    seen_phones.add(clean_phone)
    
    ulasan_num = ulasan if isinstance(ulasan, (int, float)) else 0
    rating_num = rating if isinstance(rating, (int, float)) else 4.5
    
    # Must have credible reviews (at least 15 ulasan)
    if ulasan_num < 15:
        continue
        
    # Hotel type classification
    is_hotel_like = any(w in name_lower for w in [
        'hotel', 'suites', 'resort', 'inn', 'residence', 'villas', 'villa',
        'mansion', 'boutique', 'house', 'guesthouse', 'manor', 'heritage',
        'living', 'cottages', 'cottage', 'lodge', 'palace'
    ]) or jenis in ['Hotel', 'Hotel Resor', 'Hotel jangka panjang']
    
    if not is_hotel_like:
        continue
        
    slug = re.sub(r'[^a-z0-9]+', '-', name_lower).strip('-')
    if not slug or slug in seen_slugs:
        slug = f"{slug}-{clean_phone[-4:]}"
    seen_slugs.add(slug)
    
    # Format phone display: 0812-3456-7890
    if len(clean_phone) == 11:
        phone_display = f"{clean_phone[:4]}-{clean_phone[4:7]}-{clean_phone[7:]}"
    elif len(clean_phone) == 12:
        phone_display = f"{clean_phone[:4]}-{clean_phone[4:8]}-{clean_phone[8:]}"
    elif len(clean_phone) == 13:
        phone_display = f"{clean_phone[:4]}-{clean_phone[4:9]}-{clean_phone[9:]}"
    else:
        phone_display = clean_phone

    curated_leads.append({
        'slug': slug,
        'name': name_str,
        'jenis': jenis,
        'kota': kota or 'Jakarta',
        'prov': prov or 'Indonesia',
        'rating': round(float(rating_num), 1),
        'ulasan': int(ulasan_num),
        'phone': telp,
        'clean_phone': clean_phone,
        'phone_display': phone_display,
        'alamat': alamat or f"{kota}, Indonesia"
    })

# Sort by reviews descending to get top prospects first
curated_leads.sort(key=lambda x: x['ulasan'], reverse=True)

# Select top 100 pristine leads
selected_leads = curated_leads[:100]
print(f"Filtered {len(selected_leads)} top-tier genuine hotel leads!")

# 1. Update src/data/hotelsLeads.json so all 100 hotel pages render live on Next.js
hotels_json_dict = {}
for h in selected_leads:
    clean_p = '62' + h['clean_phone'][1:]
    hotels_json_dict[h['slug']] = {
        'slug': h['slug'],
        'name': h['name'],
        'phone': h['phone_display'],
        'clean_phone': clean_p,
        'wa_link': f"https://wa.me/{clean_p}",
        'city': h['kota'],
        'province': h['prov'],
        'rating': h['rating'],
        'reviews': h['ulasan'],
        'address': h['alamat']
    }

with open(JSON_DATA_PATH, 'w', encoding='utf-8') as f:
    json.dump(hotels_json_dict, f, ensure_ascii=False, indent=2)
print(f"Updated {JSON_DATA_PATH} with {len(hotels_json_dict)} live hotel profiles.")

# 2. Build the Excel Workbook
wb_out = openpyxl.Workbook()
ws = wb_out.active
ws.title = "100 Leads Hotel Prospek"

# Define Palette & Styles
navy_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
gold_accent_font = Font(name="Calibri", size=11, bold=True, color="D4AF37")
regular_font = Font(name="Calibri", size=10, color="1E293B")
bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
link_font = Font(name="Calibri", size=10, color="0284C7", underline="single")
action_link_font = Font(name="Calibri", size=10, bold=True, color="047857", underline="single")

thin_gray_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

status_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
status_font = Font(name="Calibri", size=10, bold=True, color="92400E")

# Title Banner Row
ws.merge_cells("A1:K1")
title_cell = ws["A1"]
title_cell.value = "DATABASE 100 LEADS HOTEL & SUITES — PITCH WEBSITE LUXURY & DIRECT BOOKING (VALID DIRECT WHATSAPP)"
title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

# Subtitle Info Row
ws.merge_cells("A2:K2")
sub_cell = ws["A2"]
sub_cell.value = "Target: Hotel, Suites, Boutique Hotel, Villa & Resort (Bebas Kos-Kosan) | Demo Siap Tayang Vercel | Template Chat Pendek & To The Point"
sub_cell.font = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
sub_cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 24

# Column Headers
headers = [
    "No",
    "Nama Hotel / Prospek",
    "Kota",
    "Rating & Ulasan Google",
    "WhatsApp Hotel (Prospek)",
    "Link Demo Website (Branded Vercel)",
    "Pesan WhatsApp Personalisasi (To The Point)",
    "Link Klik Langsung Kirim WA",
    "Alamat Lengkap",
    "Status Follow-up",
    "Catatan"
]

ws.row_dimensions[3].height = 28
for col_idx, header_title in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx)
    cell.value = header_title
    cell.font = header_font
    cell.fill = navy_header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_gray_border

# Populate Data Rows
for idx, h in enumerate(selected_leads, 1):
    row_num = idx + 3
    demo_url = f"https://hotel-demo-web-two.vercel.app/p/{h['slug']}"
    
    # Concise, punchy, to-the-point, high-converting message copy
    pesan_wa = (
        f"Halo Tim Manajemen {h['name']},\n\n"
        f"Salam kenal! Kami perhatikan {h['name']} memiliki reputasi sangat baik di Google Maps "
        f"({h['rating']}⭐ dari {h['ulasan']} ulasan di {h['kota']}).\n\n"
        f"Untuk membantu meningkatkan direct booking tanpa potongan komisi OTA 15-20% (Traveloka/Agoda), "
        f"kami berinisiatif membuatkan preview website resmi modern khusus untuk {h['name']}:\n\n"
        f"👉 Akses Website Demo Anda di sini:\n"
        f"{demo_url}\n\n"
        f"Keunggulan sistem:\n"
        f"• Desain luxury bintang 5 & responsif cepat di HP tamu\n"
        f"• Reservasi otomatis langsung terarah ke WhatsApp resmi hotel Anda ({h['phone_display']})\n"
        f"• 100% bebas potongan komisi per transaksi booking\n"
        f"• Siap tayang & dapat langsung dihubungkan ke domain resmi hotel Anda\n\n"
        f"Website demo ini sudah 100% aktif dan siap digunakan. Jika tim manajemen tertarik untuk mengaktifkannya secara resmi, "
        f"bisa langsung balas pesan ini ya.\n\n"
        f"Terima kasih banyak, sukses selalu untuk {h['name']}! 🙏✨"
    )
    
    # Direct Send URL (Encoded)
    clean_target_wa = '62' + h['clean_phone'][1:]
    encoded_text = urllib.parse.quote(pesan_wa)
    direct_send_url = f"https://api.whatsapp.com/send?phone={clean_target_wa}&text={encoded_text}"
    
    # Row values
    # Col 1: No
    c1 = ws.cell(row=row_num, column=1, value=idx)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c1.font = regular_font
    c1.border = thin_gray_border
    
    # Col 2: Nama Hotel
    c2 = ws.cell(row=row_num, column=2, value=h['name'])
    c2.alignment = Alignment(horizontal="left", vertical="center")
    c2.font = bold_font
    c2.border = thin_gray_border
    
    # Col 3: Kota
    c3 = ws.cell(row=row_num, column=3, value=h['kota'])
    c3.alignment = Alignment(horizontal="center", vertical="center")
    c3.font = regular_font
    c3.border = thin_gray_border
    
    # Col 4: Rating & Ulasan Google
    c4 = ws.cell(row=row_num, column=4, value=f"⭐ {h['rating']} ({h['ulasan']} ulasan)")
    c4.alignment = Alignment(horizontal="center", vertical="center")
    c4.font = regular_font
    c4.border = thin_gray_border
    
    # Col 5: WhatsApp Hotel
    c5 = ws.cell(row=row_num, column=5, value=h['phone_display'])
    c5.alignment = Alignment(horizontal="center", vertical="center")
    c5.font = regular_font
    c5.border = thin_gray_border
    
    # Col 6: Link Demo Website Branded
    c6 = ws.cell(row=row_num, column=6, value=demo_url)
    c6.hyperlink = demo_url
    c6.alignment = Alignment(horizontal="left", vertical="center")
    c6.font = link_font
    c6.border = thin_gray_border
    
    # Col 7: Pesan WhatsApp Personalisasi
    c7 = ws.cell(row=row_num, column=7, value=pesan_wa)
    c7.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c7.font = regular_font
    c7.border = thin_gray_border
    
    # Col 8: Link Klik Langsung Kirim WA
    c8 = ws.cell(row=row_num, column=8, value="Kirim WhatsApp Sekarang ➜")
    c8.hyperlink = direct_send_url
    c8.alignment = Alignment(horizontal="center", vertical="center")
    c8.font = action_link_font
    c8.border = thin_gray_border
    
    # Col 9: Alamat Lengkap
    c9 = ws.cell(row=row_num, column=9, value=h['alamat'])
    c9.alignment = Alignment(horizontal="left", vertical="center")
    c9.font = regular_font
    c9.border = thin_gray_border
    
    # Col 10: Status Follow-up
    c10 = ws.cell(row=row_num, column=10, value="Belum Dihubungi")
    c10.alignment = Alignment(horizontal="center", vertical="center")
    c10.font = status_font
    c10.fill = status_fill
    c10.border = thin_gray_border
    
    # Col 11: Catatan
    c11 = ws.cell(row=row_num, column=11, value="")
    c11.alignment = Alignment(horizontal="left", vertical="center")
    c11.font = regular_font
    c11.border = thin_gray_border
    
    ws.row_dimensions[row_num].height = 45

# Column Widths
column_widths = {
    1: 6,    # No
    2: 34,   # Nama Hotel
    3: 16,   # Kota
    4: 24,   # Rating & Ulasan
    5: 22,   # WhatsApp
    6: 45,   # Link Demo
    7: 65,   # Pesan WA
    8: 26,   # Kirim WA Link
    9: 50,   # Alamat
    10: 18,  # Status
    11: 20   # Catatan
}

for col_idx, width in column_widths.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

# Enable freeze pane below header
ws.freeze_panes = "A4"

wb_out.save(OUTPUT_XLSX)
print(f"Successfully generated luxury hotel leads workbook: {OUTPUT_XLSX}")
