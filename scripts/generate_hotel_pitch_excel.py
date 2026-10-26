import json
import urllib.parse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_pitch_message(hotel_name, city, rating, reviews, slug, phone):
    rating_str = f"{rating:.1f}" if isinstance(rating, (int, float)) else str(rating)
    reviews_str = str(reviews)
    
    return f"""Halo Tim Manajemen {hotel_name},

Salam kenal! Saya perhatikan {hotel_name} memiliki reputasi yang sangat baik di Google Maps (⭐ {rating_str} dari {reviews_str} ulasan tamu di {city}).

Melihat antusiasme tamu dan potensi peningkatan reservasi langsung (tanpa potongan komisi OTA), kami berinisiatif membuatkan preview live website luxury hotel modern khusus untuk {hotel_name}, lengkap dengan sistem Direct WhatsApp Booking & Multi-Currency Converter:

👉 Akses Website Demo Hotel Anda di sini:
https://hotel-demo-web-two.vercel.app/p/{slug}

Keunggulan sistem website ini:
• Desain Mewah Kelas Dunia (Light Luxury Theme) & Super Responsif di Smartphone Tamu
• Direct Booking WhatsApp otomatis langsung terarah ke nomor resmi hotel Anda ({phone})
• Pemilihan Kamar & Suite Interaktif dengan konversi mata uang (USD, IDR, AUD, EUR, SGD)
• Booking Bar Real-time Check-in/Check-out & Reservasi Meja Restoran/Lounge
• Siap dihubungkan ke Domain Resmi Hotel (contoh: www.{slug}.com)

Website ini sudah live dan siap kami serahkan / aktivasi untuk manajemen hotel Anda. Jika tertarik mengklaim dan mengaktifkannya, silakan hubungi WhatsApp kami di +62 857-1820-3535 atau klik tombol 'Klaim Website Ini' di bagian atas halaman demo.

Terima kasih banyak, sukses selalu untuk okupansi dan kemajuan {hotel_name}! 🙏🏨✨"""

def main():
    with open('src/data/hotelsLeads.json', 'r', encoding='utf-8') as f:
        hotels_map = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads Hotel Pitch Prospek"

    headers = [
        "No",
        "Nama Hotel",
        "Kota / Wilayah",
        "Rating & Ulasan Google",
        "WhatsApp Hotel (Prospek)",
        "Link Demo Personal (Vercel)",
        "Pesan WhatsApp Personalisasi (Siap Kirim)",
        "Link Klik Langsung Kirim WA",
        "Alamat Lengkap",
        "Status Follow-up",
        "Catatan"
    ]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    row_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    link_font = Font(name="Segoe UI", size=10, color="0284C7", underline="single")
    wa_action_font = Font(name="Segoe UI", size=10, bold=True, color="16A34A")

    ws.row_dimensions[1].height = 36

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border

    row_idx = 2
    for hotel in hotels_map.values():
        name = hotel.get("name", "")
        city = f"{hotel.get('city', '')}, {hotel.get('province', '')}".strip(", ")
        rating = hotel.get("rating", 4.5)
        reviews = hotel.get("reviews", 0)
        phone = hotel.get("phone", "")
        clean_phone = hotel.get("clean_phone", "")
        slug = hotel.get("slug", "")
        address = hotel.get("address", "")

        rating_ulasan = f"⭐ {rating:.1f} ({reviews} ulasan)"
        demo_url = f"https://hotel-demo-web-two.vercel.app/p/{slug}"

        msg = build_pitch_message(name, hotel.get('city', ''), rating, reviews, slug, phone)
        encoded_msg = urllib.parse.quote(msg)
        wa_target = clean_phone if clean_phone else "62811388555"
        wa_direct_url = f"https://api.whatsapp.com/send?phone={wa_target}&text={encoded_msg}"

        # No
        c1 = ws.cell(row=row_idx, column=1, value=row_idx - 1)
        c1.alignment = Alignment(horizontal="center", vertical="top")

        # Nama Hotel
        c2 = ws.cell(row=row_idx, column=2, value=name)
        c2.font = bold_font
        c2.alignment = Alignment(horizontal="left", vertical="top")

        # Kota
        c3 = ws.cell(row=row_idx, column=3, value=city)
        c3.alignment = Alignment(horizontal="left", vertical="top")

        # Rating & Ulasan
        c4 = ws.cell(row=row_idx, column=4, value=rating_ulasan)
        c4.alignment = Alignment(horizontal="center", vertical="top")

        # Phone
        c5 = ws.cell(row=row_idx, column=5, value=phone)
        c5.alignment = Alignment(horizontal="center", vertical="top")

        # Demo URL
        c6 = ws.cell(row=row_idx, column=6)
        c6.value = f'=HYPERLINK("{demo_url}", "{demo_url}")'
        c6.font = link_font
        c6.alignment = Alignment(horizontal="left", vertical="top")

        # Pesan WhatsApp
        c7 = ws.cell(row=row_idx, column=7, value=msg)
        c7.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Link Klik Langsung Kirim WA
        c8 = ws.cell(row=row_idx, column=8)
        c8.value = f'=HYPERLINK("{wa_direct_url}", "Chat WA Sekarang ->")'
        c8.font = wa_action_font
        c8.alignment = Alignment(horizontal="center", vertical="top")

        # Alamat
        c9 = ws.cell(row=row_idx, column=9, value=address)
        c9.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Status
        c10 = ws.cell(row=row_idx, column=10, value="Belum Dihubungi")
        c10.alignment = Alignment(horizontal="center", vertical="top")

        # Catatan
        c11 = ws.cell(row=row_idx, column=11, value="")
        c11.alignment = Alignment(horizontal="left", vertical="top")

        for col_i in range(1, 12):
            cell_i = ws.cell(row=row_idx, column=col_i)
            cell_i.border = cell_border
            if col_i not in (2, 6, 8):
                cell_i.font = row_font

        ws.row_dimensions[row_idx].height = 48
        row_idx += 1

    col_widths = {
        1: 6.0,
        2: 32.0,
        3: 20.0,
        4: 20.0,
        5: 18.0,
        6: 38.0,
        7: 65.0,
        8: 22.0,
        9: 38.0,
        10: 16.0,
        11: 20.0
    }
    for col_i, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_i)].width = width

    ws.freeze_panes = "C2"

    output_filename = "HOTEL_PITCH_PROSPEK_LEADS.xlsx"
    wb.save(output_filename)
    print(f"Successfully generated {output_filename} with {row_idx - 2} leads!")

if __name__ == "__main__":
    main()
